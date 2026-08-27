import os
import csv
import io
from werkzeug.utils import secure_filename
from sqlalchemy import or_
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from flask import Flask, render_template, request, redirect, url_for, Response, flash
from models import (
    db, Settings, Plant, Customer, Lead, Contract, ProductionBatch,
    BusUnit, BatteryBatch, CityFleet, DailySummary, ViolationRecord, SystemUser,
    Employee, Shipment
)

BASE_DIR = os.path.abspath(os.path.dirname(__file__))
DB_PATH = os.path.join(BASE_DIR, "database", "nevf.db")
UPLOAD_FOLDER = os.path.join(BASE_DIR, "static", "uploads", "employees")
ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg"}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


app = Flask(__name__)
app.config["SQLALCHEMY_DATABASE_URI"] = f"sqlite:///{DB_PATH}"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.secret_key = "nevf-local-dev-key"

db.init_app(app)


def get_settings():
    s = Settings.query.first()
    if not s:
        s = Settings(company_name="NEVF", manufacturing_target=8000, currency="PKR")
        db.session.add(s)
        db.session.commit()
    return s


def csv_response(rows, headers, filename):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)
    return Response(buf.getvalue(), mimetype="text/csv",
                     headers={"Content-Disposition": f"attachment;filename={filename}"})


def excel_response(headers, rows, filename, title="NEVF Report"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    header_fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    title_font = Font(color="0D47A1", bold=True, size=15)
    thin = Side(border_style="thin", color="D9E0E8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = title_font
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 6

    header_row = 3
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = border
    ws.row_dimensions[header_row].height = 22

    for r_idx, row in enumerate(rows, start=header_row + 1):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = border
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if isinstance(value, bool):
                pass
            elif isinstance(value, int):
                cell.number_format = "#,##0"
            elif isinstance(value, float):
                cell.number_format = "#,##0.00"
            if r_idx % 2 == 0:
                cell.fill = PatternFill(start_color="F6F8FB", end_color="F6F8FB", fill_type="solid")

    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows:
            val = row[col_idx - 1]
            max_len = max(max_len, len(str(val)) if val is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 45)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment;filename={filename}"},
    )


# ---------------- DASHBOARD ----------------
@app.route("/")
def dashboard():
    settings = get_settings()
    total_manufactured = db.session.query(db.func.sum(ProductionBatch.units)).scalar() or 0
    target = settings.manufacturing_target or 8000
    progress_pct = round((total_manufactured / target) * 100, 1) if target else 0

    total_customers = Customer.query.count()
    govt_customers = Customer.query.filter_by(type="Government").count()
    dealer_customers = Customer.query.filter(Customer.type.in_(["Dealer", "Shop"])).count()

    daily = DailySummary.query.order_by(DailySummary.id).all()
    violations = ViolationRecord.query.order_by(ViolationRecord.id).all()
    cities = CityFleet.query.order_by(CityFleet.delayed_trips.desc()).limit(5).all()

    total_fleet_buses = db.session.query(db.func.sum(CityFleet.buses)).scalar() or 0
    total_routes = db.session.query(db.func.sum(CityFleet.routes)).scalar() or 0
    total_cities = CityFleet.query.count()

    total_violation_count = sum(v.total for v in violations)
    total_penalty = sum(v.penalty_km for v in violations)
    v_smoking = sum(v.smoking for v in violations)
    v_calling = sum(v.calling for v in violations)
    v_seatbelt = sum(v.seatbelt for v in violations)
    v_eating = sum(v.eating for v in violations)
    v_uniform = sum(v.uniform for v in violations)
    violation_breakdown = [
        ("Calling While Driving", v_calling),
        ("Not Wearing Seat Belt", v_seatbelt),
        ("Smoking While Driving", v_smoking),
        ("Eating / Drinking", v_eating),
        ("Not in Uniform", v_uniform),
    ]
    max_violation = max([x[1] for x in violation_breakdown]) or 1

    total_employees = Employee.query.count()
    total_shipments = Shipment.query.filter(Shipment.status != "Delivered").count()

    return render_template(
        "dashboard.html", active="dashboard", settings=settings,
        total_manufactured=total_manufactured, target=target, progress_pct=progress_pct,
        total_customers=total_customers, govt_customers=govt_customers, dealer_customers=dealer_customers,
        daily=daily, cities=cities, total_fleet_buses=total_fleet_buses, total_routes=total_routes,
        total_cities=total_cities, total_violation_count=total_violation_count, total_penalty=total_penalty,
        violation_breakdown=violation_breakdown, max_violation=max_violation,
        total_employees=total_employees, total_shipments=total_shipments,
    )


# ---------------- PLANTS ----------------
@app.route("/plants")
def plants():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "desc")

    sort_columns = {
        "name": Plant.name, "city": Plant.city, "lines": Plant.lines,
        "capacity_per_month": Plant.capacity_per_month, "status": Plant.status, "id": Plant.id,
    }
    col = sort_columns.get(sort, Plant.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = Plant.query.order_by(col).all()

    status_counts = {}
    for p in rows:
        status_counts[p.status] = status_counts.get(p.status, 0) + 1
    status_breakdown = sorted(status_counts.items(), key=lambda x: -x[1])
    max_status = max([x[1] for x in status_breakdown]) if status_breakdown else 1

    total_capacity = sum(p.capacity_per_month for p in rows)

    return render_template("plants.html", active="plants", plants=rows,
                            current_sort=sort, current_dir=direction,
                            status_breakdown=status_breakdown, max_status=max_status,
                            total_capacity=total_capacity)


@app.route("/plants/add", methods=["POST"])
def add_plant():
    p = Plant(
        name=request.form["name"], city=request.form["city"],
        address=request.form.get("address", ""),
        capacity_per_month=int(request.form.get("capacity_per_month") or 0),
        lines=int(request.form.get("lines") or 1),
        status=request.form.get("status", "Operational"),
        established_date=request.form.get("established_date", ""),
        notes=request.form.get("notes", ""),
    )
    db.session.add(p)
    db.session.commit()
    flash("Plant added successfully.")
    return redirect(url_for("plants"))


@app.route("/plants/edit/<int:id>", methods=["POST"])
def edit_plant(id):
    p = Plant.query.get_or_404(id)
    p.name = request.form["name"]
    p.city = request.form["city"]
    p.address = request.form.get("address", "")
    p.capacity_per_month = int(request.form.get("capacity_per_month") or 0)
    p.lines = int(request.form.get("lines") or 1)
    p.status = request.form.get("status", "Operational")
    p.established_date = request.form.get("established_date", "")
    p.notes = request.form.get("notes", "")
    db.session.commit()
    flash("Plant updated successfully.")
    return redirect(url_for("plants"))


@app.route("/plants/delete/<int:id>", methods=["POST"])
def delete_plant(id):
    p = Plant.query.get_or_404(id)
    db.session.delete(p)
    db.session.commit()
    return redirect(url_for("plants"))


@app.route("/plants/export")
def export_plants():
    rows = Plant.query.all()
    data = [[p.name, p.city, p.lines, p.capacity_per_month, p.status, p.established_date] for p in rows]
    return excel_response(
        ["Plant Name", "City", "Lines", "Capacity/Month", "Status", "Established"],
        data, "plants.xlsx", title="NEVF — Manufacturing Plants Report"
    )


# ---------------- CUSTOMERS ----------------
@app.route("/customers")
def customers():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "desc")

    sort_columns = {
        "name": Customer.name, "type": Customer.type, "contact_person": Customer.contact_person,
        "phone": Customer.phone, "city": Customer.city, "status": Customer.status, "id": Customer.id,
    }
    col = sort_columns.get(sort, Customer.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = Customer.query.order_by(col).all()

    type_counts = {}
    city_counts = {}
    for c in rows:
        type_counts[c.type] = type_counts.get(c.type, 0) + 1
        if c.city:
            city_counts[c.city] = city_counts.get(c.city, 0) + 1
    type_breakdown = sorted(type_counts.items(), key=lambda x: -x[1])
    city_breakdown = sorted(city_counts.items(), key=lambda x: -x[1])[:8]
    max_type = max([x[1] for x in type_breakdown]) if type_breakdown else 1
    max_city = max([x[1] for x in city_breakdown]) if city_breakdown else 1

    return render_template(
        "customers.html", active="customers", customers=rows, total=len(rows),
        govt=len([c for c in rows if c.type == "Government"]),
        dealer=len([c for c in rows if c.type in ("Dealer", "Shop")]),
        type_breakdown=type_breakdown, city_breakdown=city_breakdown,
        max_type=max_type, max_city=max_city,
        current_sort=sort, current_dir=direction,
    )


@app.route("/customers/add", methods=["POST"])
def add_customer():
    c = Customer(
        name=request.form["name"], type=request.form["type"],
        contact_person=request.form.get("contact_person", ""), phone=request.form.get("phone", ""),
        email=request.form.get("email", ""), city=request.form.get("city", ""),
        status=request.form.get("status", "Active"), notes=request.form.get("notes", ""),
    )
    db.session.add(c)
    db.session.commit()
    flash("Customer added successfully.")
    return redirect(url_for("customers"))


@app.route("/customers/edit/<int:id>", methods=["POST"])
def edit_customer(id):
    c = Customer.query.get_or_404(id)
    c.name = request.form["name"]
    c.type = request.form["type"]
    c.contact_person = request.form.get("contact_person", "")
    c.phone = request.form.get("phone", "")
    c.email = request.form.get("email", "")
    c.city = request.form.get("city", "")
    c.status = request.form.get("status", "Active")
    c.notes = request.form.get("notes", "")
    db.session.commit()
    flash("Customer updated successfully.")
    return redirect(url_for("customers"))


@app.route("/customers/delete/<int:id>", methods=["POST"])
def delete_customer(id):
    c = Customer.query.get_or_404(id)
    db.session.delete(c)
    db.session.commit()
    return redirect(url_for("customers"))


@app.route("/customers/export")
def export_customers():
    rows = Customer.query.all()
    data = [[c.name, c.type, c.contact_person, c.phone, c.city, c.status] for c in rows]
    return excel_response(
        ["Name", "Type", "Contact Person", "Phone", "City", "Status"],
        data, "customers.xlsx", title="NEVF — Customers Report"
    )


# ---------------- LEADS ----------------
@app.route("/leads")
def leads():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "desc")

    sort_columns = {
        "title": Lead.title, "lead_type": Lead.lead_type, "value": Lead.value,
        "stage": Lead.stage, "expected_date": Lead.expected_date, "id": Lead.id,
    }
    col = sort_columns.get(sort, Lead.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = Lead.query.order_by(col).all()

    stage_counts = {}
    for l in rows:
        stage_counts[l.stage] = stage_counts.get(l.stage, 0) + 1
    stage_breakdown = sorted(stage_counts.items(), key=lambda x: -x[1])
    max_stage = max([x[1] for x in stage_breakdown]) if stage_breakdown else 1

    total_value = sum(l.value for l in rows)
    won_value = sum(l.value for l in rows if l.stage == "Won")

    return render_template(
        "leads.html", active="leads", leads=rows, current_sort=sort, current_dir=direction,
        stage_breakdown=stage_breakdown, max_stage=max_stage,
        total_value=total_value, won_value=won_value,
    )


@app.route("/leads/add", methods=["POST"])
def add_lead():
    l = Lead(
        title=request.form["title"], lead_type=request.form.get("lead_type", ""),
        value=float(request.form.get("value") or 0), stage=request.form.get("stage", "New"),
        expected_date=request.form.get("expected_date", ""), notes=request.form.get("notes", ""),
    )
    db.session.add(l)
    db.session.commit()
    flash("Lead added successfully.")
    return redirect(url_for("leads"))


@app.route("/leads/edit/<int:id>", methods=["POST"])
def edit_lead(id):
    l = Lead.query.get_or_404(id)
    l.title = request.form["title"]
    l.lead_type = request.form.get("lead_type", "")
    l.value = float(request.form.get("value") or 0)
    l.stage = request.form.get("stage", "New")
    l.expected_date = request.form.get("expected_date", "")
    l.notes = request.form.get("notes", "")
    db.session.commit()
    flash("Lead updated successfully.")
    return redirect(url_for("leads"))


@app.route("/leads/delete/<int:id>", methods=["POST"])
def delete_lead(id):
    l = Lead.query.get_or_404(id)
    db.session.delete(l)
    db.session.commit()
    return redirect(url_for("leads"))


@app.route("/leads/export")
def export_leads():
    rows = Lead.query.all()
    data = [[l.title, l.lead_type, l.value, l.stage, l.expected_date] for l in rows]
    return excel_response(
        ["Title", "Type", "Estimated Value", "Stage", "Expected Date"],
        data, "leads.xlsx", title="NEVF — Leads Report"
    )


# ---------------- CONTRACTS ----------------
@app.route("/contracts")
def contracts():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "desc")

    sort_columns = {
        "ref": Contract.ref, "customer_name": Contract.customer_name, "item": Contract.item,
        "qty": Contract.qty, "value": Contract.value, "status": Contract.status, "id": Contract.id,
    }
    col = sort_columns.get(sort, Contract.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = Contract.query.order_by(col).all()

    status_counts = {}
    for c in rows:
        status_counts[c.status] = status_counts.get(c.status, 0) + 1
    status_breakdown = sorted(status_counts.items(), key=lambda x: -x[1])
    max_status = max([x[1] for x in status_breakdown]) if status_breakdown else 1

    total_value = sum(c.value for c in rows)
    delivered_value = sum(c.value for c in rows if c.status == "Delivered")

    return render_template(
        "contracts.html", active="contracts", contracts=rows, current_sort=sort, current_dir=direction,
        status_breakdown=status_breakdown, max_status=max_status,
        total_value=total_value, delivered_value=delivered_value,
    )


@app.route("/contracts/add", methods=["POST"])
def add_contract():
    c = Contract(
        ref=request.form["ref"], customer_name=request.form.get("customer_name", ""),
        item=request.form.get("item", "Electric Bus"), qty=int(request.form.get("qty") or 0),
        value=float(request.form.get("value") or 0), status=request.form.get("status", "Draft"),
    )
    db.session.add(c)
    db.session.commit()
    flash("Contract added successfully.")
    return redirect(url_for("contracts"))


@app.route("/contracts/edit/<int:id>", methods=["POST"])
def edit_contract(id):
    c = Contract.query.get_or_404(id)
    c.ref = request.form["ref"]
    c.customer_name = request.form.get("customer_name", "")
    c.item = request.form.get("item", "Electric Bus")
    c.qty = int(request.form.get("qty") or 0)
    c.value = float(request.form.get("value") or 0)
    c.status = request.form.get("status", "Draft")
    db.session.commit()
    flash("Contract updated successfully.")
    return redirect(url_for("contracts"))


@app.route("/contracts/delete/<int:id>", methods=["POST"])
def delete_contract(id):
    c = Contract.query.get_or_404(id)
    db.session.delete(c)
    db.session.commit()
    return redirect(url_for("contracts"))


@app.route("/contracts/export")
def export_contracts():
    rows = Contract.query.all()
    data = [[c.ref, c.customer_name, c.item, c.qty, c.value, c.status] for c in rows]
    return excel_response(
        ["Contract Ref", "Customer", "Item", "Qty", "Value (PKR)", "Status"],
        data, "contracts.xlsx", title="NEVF — Contracts & Orders Report"
    )


# ---------------- PRODUCTION ----------------
@app.route("/production")
def production():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "desc")

    sort_columns = {
        "batch_id": ProductionBatch.batch_id, "model": ProductionBatch.model,
        "units": ProductionBatch.units, "date": ProductionBatch.date,
        "qc_status": ProductionBatch.qc_status, "id": ProductionBatch.id,
    }
    col = sort_columns.get(sort, ProductionBatch.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = ProductionBatch.query.order_by(col).all()
    settings = get_settings()
    total_units = db.session.query(db.func.sum(ProductionBatch.units)).scalar() or 0
    target = settings.manufacturing_target or 8000
    pct = round((total_units / target) * 100, 1) if target else 0
    plant_list = Plant.query.all()

    qc_counts = {}
    for b in rows:
        qc_counts[b.qc_status] = qc_counts.get(b.qc_status, 0) + 1
    qc_breakdown = sorted(qc_counts.items(), key=lambda x: -x[1])
    max_qc = max([x[1] for x in qc_breakdown]) if qc_breakdown else 1

    return render_template("production.html", active="production", batches=rows, total_units=total_units,
                            target=target, pct=pct, plant_list=plant_list,
                            current_sort=sort, current_dir=direction,
                            qc_breakdown=qc_breakdown, max_qc=max_qc)


@app.route("/production/add", methods=["POST"])
def add_production():
    b = ProductionBatch(
        batch_id=request.form["batch_id"], model=request.form.get("model", ""),
        units=int(request.form.get("units") or 0), plant_id=request.form.get("plant_id") or None,
        date=request.form.get("date", ""), qc_status=request.form.get("qc_status", "Passed"),
    )
    db.session.add(b)
    db.session.commit()
    flash("Production batch logged.")
    return redirect(url_for("production"))


@app.route("/production/edit/<int:id>", methods=["POST"])
def edit_production(id):
    b = ProductionBatch.query.get_or_404(id)
    b.batch_id = request.form["batch_id"]
    b.model = request.form.get("model", "")
    b.units = int(request.form.get("units") or 0)
    b.plant_id = request.form.get("plant_id") or None
    b.date = request.form.get("date", "")
    b.qc_status = request.form.get("qc_status", "Passed")
    db.session.commit()
    flash("Production batch updated successfully.")
    return redirect(url_for("production"))


@app.route("/production/delete/<int:id>", methods=["POST"])
def delete_production(id):
    b = ProductionBatch.query.get_or_404(id)
    db.session.delete(b)
    db.session.commit()
    return redirect(url_for("production"))


@app.route("/production/export")
def export_production():
    rows = ProductionBatch.query.all()
    data = [[b.batch_id, b.model, b.units, b.plant.name if b.plant else "-", b.date, b.qc_status] for b in rows]
    return excel_response(
        ["Batch ID", "Model", "Units Produced", "Plant", "Date", "QC Status"],
        data, "production_batches.xlsx", title="NEVF — Production Batches Report"
    )


# ---------------- INVENTORY: BUSES ----------------
@app.route("/inventory/buses")
def inventory_buses():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "desc")

    sort_columns = {
        "unit_no": BusUnit.unit_no, "model": BusUnit.model,
        "status": BusUnit.status, "location": BusUnit.location, "id": BusUnit.id,
    }
    col = sort_columns.get(sort, BusUnit.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = BusUnit.query.order_by(col).all()

    status_counts = {}
    for b in rows:
        status_counts[b.status] = status_counts.get(b.status, 0) + 1
    status_breakdown = sorted(status_counts.items(), key=lambda x: -x[1])
    max_status = max([x[1] for x in status_breakdown]) if status_breakdown else 1

    return render_template("inventory_buses.html", active="inv_buses", buses=rows,
                            in_stock=len([b for b in rows if b.status == "In Stock"]),
                            dispatched=len([b for b in rows if b.status == "Dispatched"]),
                            under_qc=len([b for b in rows if b.status == "Under QC"]),
                            current_sort=sort, current_dir=direction,
                            status_breakdown=status_breakdown, max_status=max_status)


@app.route("/inventory/buses/add", methods=["POST"])
def add_bus():
    b = BusUnit(
        unit_no=request.form["unit_no"], model=request.form.get("model", ""),
        battery_fitted=request.form.get("battery_fitted", "No"),
        status=request.form.get("status", "In Stock"), location=request.form.get("location", ""),
    )
    db.session.add(b)
    db.session.commit()
    flash("Bus unit added.")
    return redirect(url_for("inventory_buses"))


@app.route("/inventory/buses/edit/<int:id>", methods=["POST"])
def edit_bus(id):
    b = BusUnit.query.get_or_404(id)
    b.unit_no = request.form["unit_no"]
    b.model = request.form.get("model", "")
    b.battery_fitted = request.form.get("battery_fitted", "No")
    b.status = request.form.get("status", "In Stock")
    b.location = request.form.get("location", "")
    db.session.commit()
    flash("Bus unit updated successfully.")
    return redirect(url_for("inventory_buses"))


@app.route("/inventory/buses/delete/<int:id>", methods=["POST"])
def delete_bus(id):
    b = BusUnit.query.get_or_404(id)
    db.session.delete(b)
    db.session.commit()
    return redirect(url_for("inventory_buses"))


@app.route("/inventory/buses/export")
def export_buses():
    rows = BusUnit.query.all()
    data = [[b.unit_no, b.model, b.battery_fitted, b.status, b.location] for b in rows]
    return excel_response(
        ["Chassis / Unit No.", "Model", "Battery Fitted", "Status", "Location"],
        data, "buses_stock.xlsx", title="NEVF — Buses Stock Report"
    )


# ---------------- INVENTORY: BATTERIES ----------------
@app.route("/inventory/batteries")
def inventory_batteries():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "desc")

    sort_columns = {
        "batch": BatteryBatch.batch, "capacity_kwh": BatteryBatch.capacity_kwh,
        "supplier": BatteryBatch.supplier, "qty": BatteryBatch.qty,
        "status": BatteryBatch.status, "id": BatteryBatch.id,
    }
    col = sort_columns.get(sort, BatteryBatch.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = BatteryBatch.query.order_by(col).all()

    status_counts = {}
    for b in rows:
        status_counts[b.status] = status_counts.get(b.status, 0) + b.qty
    status_breakdown = sorted(status_counts.items(), key=lambda x: -x[1])
    max_status = max([x[1] for x in status_breakdown]) if status_breakdown else 1

    return render_template("inventory_batteries.html", active="inv_batteries", batteries=rows,
                            total_units=sum(b.qty for b in rows),
                            available=sum(b.qty for b in rows if b.status == "Available"),
                            fitted=sum(b.qty for b in rows if b.status == "Fitted to Buses"),
                            current_sort=sort, current_dir=direction,
                            status_breakdown=status_breakdown, max_status=max_status)


@app.route("/inventory/batteries/add", methods=["POST"])
def add_battery():
    b = BatteryBatch(
        batch=request.form["batch"], capacity_kwh=float(request.form.get("capacity_kwh") or 0),
        supplier=request.form.get("supplier", ""), qty=int(request.form.get("qty") or 0),
        status=request.form.get("status", "Available"),
    )
    db.session.add(b)
    db.session.commit()
    flash("Battery batch added.")
    return redirect(url_for("inventory_batteries"))


@app.route("/inventory/batteries/edit/<int:id>", methods=["POST"])
def edit_battery(id):
    b = BatteryBatch.query.get_or_404(id)
    b.batch = request.form["batch"]
    b.capacity_kwh = float(request.form.get("capacity_kwh") or 0)
    b.supplier = request.form.get("supplier", "")
    b.qty = int(request.form.get("qty") or 0)
    b.status = request.form.get("status", "Available")
    db.session.commit()
    flash("Battery batch updated successfully.")
    return redirect(url_for("inventory_batteries"))


@app.route("/inventory/batteries/delete/<int:id>", methods=["POST"])
def delete_battery(id):
    b = BatteryBatch.query.get_or_404(id)
    db.session.delete(b)
    db.session.commit()
    return redirect(url_for("inventory_batteries"))


@app.route("/inventory/batteries/export")
def export_batteries():
    rows = BatteryBatch.query.all()
    data = [[b.batch, b.capacity_kwh, b.supplier, b.qty, b.status] for b in rows]
    return excel_response(
        ["Batch / Serial", "Capacity (kWh)", "Supplier", "Qty", "Status"],
        data, "batteries_stock.xlsx", title="NEVF — Batteries Stock Report"
    )


# ---------------- FLEET: CITY SUMMARY (with sorting) ----------------
@app.route("/fleet/city")
def fleet_city():
    sort = request.args.get("sort", "city")
    direction = request.args.get("dir", "asc")

    sort_columns = {
        "city": CityFleet.city, "buses": CityFleet.buses, "routes": CityFleet.routes,
        "planned": CityFleet.planned_trips, "matured": CityFleet.matured_trips,
        "delayed": CityFleet.delayed_trips, "delay_pct": CityFleet.delay_pct,
    }
    col = sort_columns.get(sort, CityFleet.city)
    col = col.desc() if direction == "desc" else col.asc()

    rows = CityFleet.query.order_by(col).all()
    totals = {"buses": sum(c.buses for c in rows), "routes": sum(c.routes for c in rows),
              "planned": sum(c.planned_trips for c in rows), "delayed": sum(c.delayed_trips for c in rows)}
    overall_delay = round((totals["delayed"] / totals["planned"]) * 100, 1) if totals["planned"] else 0
    top_delay_cities = sorted(rows, key=lambda c: -c.delay_pct)[:10]
    return render_template("fleet_city.html", active="fleet_city", cities=rows, totals=totals,
                            overall_delay=overall_delay, current_sort=sort, current_dir=direction,
                            top_delay_cities=top_delay_cities)


@app.route("/fleet/city/export")
def export_city():
    rows = CityFleet.query.all()
    data = [[c.city, c.buses, c.routes, c.planned_trips, c.matured_trips, c.delayed_trips, c.delay_pct] for c in rows]
    return excel_response(
        ["City", "Buses", "Routes", "Planned Trips", "Matured Trips", "Delayed Trips", "Delay %"],
        data, "city_summary.xlsx", title="NEVF — City Summary Report"
    )


# ---------------- FLEET: DAILY SUMMARY (with sorting) ----------------
@app.route("/fleet/daily")
def fleet_daily():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "asc")

    sort_columns = {
        "date": DailySummary.date_label, "buses": DailySummary.buses, "routes": DailySummary.routes,
        "planned": DailySummary.planned_trips, "matured": DailySummary.matured_trips,
        "delayed": DailySummary.delayed_trips, "delay_pct": DailySummary.delay_pct,
        "violations": DailySummary.violations, "penalty": DailySummary.penalty_km, "id": DailySummary.id,
    }
    col = sort_columns.get(sort, DailySummary.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = DailySummary.query.order_by(col).all()
    max_delay = max([d.delay_pct for d in rows]) if rows else 1
    return render_template("fleet_daily.html", active="fleet_daily", daily=rows, max_delay=max_delay,
                            current_sort=sort, current_dir=direction)


@app.route("/fleet/daily/add", methods=["POST"])
def add_daily_summary():
    d = DailySummary(
        date_label=request.form["date_label"],
        buses=int(request.form.get("buses") or 0),
        routes=int(request.form.get("routes") or 0),
        planned_trips=int(request.form.get("planned_trips") or 0),
        matured_trips=int(request.form.get("matured_trips") or 0),
        delayed_trips=int(request.form.get("delayed_trips") or 0),
        delay_pct=float(request.form.get("delay_pct") or 0),
        violations=int(request.form.get("violations") or 0),
        penalty_km=int(request.form.get("penalty_km") or 0),
    )
    db.session.add(d)
    db.session.commit()
    flash("Daily entry added successfully.")
    return redirect(url_for("fleet_daily"))


@app.route("/fleet/daily/edit/<int:id>", methods=["POST"])
def edit_daily_summary(id):
    d = DailySummary.query.get_or_404(id)
    d.date_label = request.form["date_label"]
    d.buses = int(request.form.get("buses") or 0)
    d.routes = int(request.form.get("routes") or 0)
    d.planned_trips = int(request.form.get("planned_trips") or 0)
    d.matured_trips = int(request.form.get("matured_trips") or 0)
    d.delayed_trips = int(request.form.get("delayed_trips") or 0)
    d.delay_pct = float(request.form.get("delay_pct") or 0)
    d.violations = int(request.form.get("violations") or 0)
    d.penalty_km = int(request.form.get("penalty_km") or 0)
    db.session.commit()
    flash("Daily entry updated successfully.")
    return redirect(url_for("fleet_daily"))


@app.route("/fleet/daily/delete/<int:id>", methods=["POST"])
def delete_daily_summary(id):
    d = DailySummary.query.get_or_404(id)
    db.session.delete(d)
    db.session.commit()
    return redirect(url_for("fleet_daily"))


@app.route("/fleet/daily/export")
def export_daily():
    rows = DailySummary.query.all()
    data = [[d.date_label, d.buses, d.routes, d.planned_trips, d.matured_trips, d.delayed_trips,
             d.delay_pct, d.violations, d.penalty_km] for d in rows]
    return excel_response(
        ["Date", "Buses", "Routes", "Planned", "Matured", "Delayed", "Delay %", "Violations", "Penalty (KM)"],
        data, "daily_summary.xlsx", title="NEVF — Daily Summary Report"
    )


# ---------------- FLEET: VIOLATIONS ----------------
@app.route("/fleet/violations")
def fleet_violations():
    sort = request.args.get("sort", "id")
    direction = request.args.get("dir", "asc")

    sort_columns = {
        "date": ViolationRecord.date_label, "smoking": ViolationRecord.smoking,
        "calling": ViolationRecord.calling, "seatbelt": ViolationRecord.seatbelt,
        "eating": ViolationRecord.eating, "uniform": ViolationRecord.uniform,
        "total": ViolationRecord.total, "penalty": ViolationRecord.penalty_km, "id": ViolationRecord.id,
    }
    col = sort_columns.get(sort, ViolationRecord.id)
    col = col.desc() if direction == "desc" else col.asc()

    rows = ViolationRecord.query.order_by(col).all()
    total_violations = sum(v.total for v in rows)
    total_penalty = sum(v.penalty_km for v in rows)
    breakdown = [
        ("Calling While Driving", sum(v.calling for v in rows)),
        ("Not Wearing Seat Belt", sum(v.seatbelt for v in rows)),
        ("Smoking While Driving", sum(v.smoking for v in rows)),
        ("Eating / Drinking", sum(v.eating for v in rows)),
        ("Not in Uniform", sum(v.uniform for v in rows)),
    ]
    max_val = max([x[1] for x in breakdown]) or 1
    return render_template("fleet_violations.html", active="fleet_violations", violations=rows,
                            total_violations=total_violations, total_penalty=total_penalty,
                            breakdown=breakdown, max_val=max_val, days_tracked=len(rows),
                            current_sort=sort, current_dir=direction)


@app.route("/fleet/violations/add", methods=["POST"])
def add_violation():
    smoking = int(request.form.get("smoking") or 0)
    calling = int(request.form.get("calling") or 0)
    seatbelt = int(request.form.get("seatbelt") or 0)
    eating = int(request.form.get("eating") or 0)
    uniform = int(request.form.get("uniform") or 0)
    total = smoking + calling + seatbelt + eating + uniform
    v = ViolationRecord(
        date_label=request.form["date_label"],
        smoking=smoking, calling=calling, seatbelt=seatbelt,
        eating=eating, uniform=uniform, total=total,
        penalty_km=int(request.form.get("penalty_km") or 0),
    )
    db.session.add(v)
    db.session.commit()
    flash("Violation entry added successfully.")
    return redirect(url_for("fleet_violations"))


@app.route("/fleet/violations/edit/<int:id>", methods=["POST"])
def edit_violation(id):
    v = ViolationRecord.query.get_or_404(id)
    v.date_label = request.form["date_label"]
    v.smoking = int(request.form.get("smoking") or 0)
    v.calling = int(request.form.get("calling") or 0)
    v.seatbelt = int(request.form.get("seatbelt") or 0)
    v.eating = int(request.form.get("eating") or 0)
    v.uniform = int(request.form.get("uniform") or 0)
    v.total = v.smoking + v.calling + v.seatbelt + v.eating + v.uniform
    v.penalty_km = int(request.form.get("penalty_km") or 0)
    db.session.commit()
    flash("Violation entry updated successfully.")
    return redirect(url_for("fleet_violations"))


@app.route("/fleet/violations/delete/<int:id>", methods=["POST"])
def delete_violation(id):
    v = ViolationRecord.query.get_or_404(id)
    db.session.delete(v)
    db.session.commit()
    return redirect(url_for("fleet_violations"))


@app.route("/fleet/violations/export")
def export_violations():
    rows = ViolationRecord.query.all()
    data = [[v.date_label, v.smoking, v.calling, v.seatbelt, v.eating, v.uniform, v.total, v.penalty_km] for v in rows]
    return excel_response(
        ["Date", "Smoking", "Calling", "Seat Belt", "Eating/Drinking", "Uniform", "Total", "Penalty (KM)"],
        data, "violations.xlsx", title="NEVF — Driver Violations Report"
    )


# ---------------- EMPLOYEES ----------------
@app.route("/employees")
def employees():
    rows = Employee.query.order_by(Employee.id).all()

    dept_counts = {}
    for e in rows:
        if e.department:
            dept_counts[e.department] = dept_counts.get(e.department, 0) + 1
    dept_breakdown = sorted(dept_counts.items(), key=lambda x: -x[1])
    max_dept = max([x[1] for x in dept_breakdown]) if dept_breakdown else 1

    top_level = [e for e in rows if e.reports_to_id is None]

    return render_template(
        "employees.html", active="employees", employees=rows, total=len(rows),
        active_count=len([e for e in rows if e.status == "Active"]),
        dept_breakdown=dept_breakdown, max_dept=max_dept, top_level=top_level,
    )


@app.route("/employees/add", methods=["POST"])
def add_employee():
    e = Employee(
        name=request.form["name"],
        designation=request.form.get("designation", ""),
        department=request.form.get("department", ""),
        phone=request.form.get("phone", ""),
        email=request.form.get("email", ""),
        status=request.form.get("status", "Active"),
        joining_date=request.form.get("joining_date", ""),
        plant_id=request.form.get("plant_id") or None,
        reports_to_id=request.form.get("reports_to_id") or None,
    )
    db.session.add(e)
    db.session.commit()

    photo = request.files.get("photo")
    if photo and photo.filename and allowed_file(photo.filename):
        ext = photo.filename.rsplit(".", 1)[1].lower()
        filename = f"emp_{e.id}.{ext}"
        os.makedirs(UPLOAD_FOLDER, exist_ok=True)
        photo.save(os.path.join(UPLOAD_FOLDER, filename))
        e.photo_filename = filename
        db.session.commit()

    flash("Employee added successfully.")
    return redirect(url_for("employees"))


@app.route("/employees/edit/<int:id>", methods=["POST"])
def edit_employee(id):
    e = Employee.query.get_or_404(id)
    e.name = request.form["name"]
    e.designation = request.form.get("designation", "")
    e.department = request.form.get("department", "")
    e.phone = request.form.get("phone", "")
    e.email = request.form.get("email", "")
    e.status = request.form.get("status", "Active")
    e.joining_date = request.form.get("joining_date", "")
    e.plant_id = request.form.get("plant_id") or None
    reports_to = request.form.get("reports_to_id") or None
    e.reports_to_id = int(reports_to) if reports_to and int(reports_to) != e.id else None

    photo = request.files.get("photo")
    if photo and photo.filename and allowed_file(photo.filename):
        ext = photo.filename.rsplit(".", 1)[1].lower()
        filename = f"emp_{e.id}.{ext}"
        os.makedirs(UPLOAD_FOLDER, exist_ok=True)
        photo.save(os.path.join(UPLOAD_FOLDER, filename))
        e.photo_filename = filename

    db.session.commit()
    flash("Employee updated successfully.")
    return redirect(url_for("employees"))


@app.route("/employees/delete/<int:id>", methods=["POST"])
def delete_employee(id):
    e = Employee.query.get_or_404(id)
    db.session.delete(e)
    db.session.commit()
    return redirect(url_for("employees"))
   
   
@app.route("/employees/export")
def export_employees():
    rows = Employee.query.all()
    data = [[e.name, e.designation, e.department, e.manager.name if e.manager else "-",
             e.phone, e.status, e.joining_date] for e in rows]
    return excel_response(
        ["Name", "Designation", "Department", "Reports To", "Phone", "Status", "Joining Date"],
        data, "employees.xlsx", title="NEVF — Employees Report"
    )


# ---------------- SHIPMENTS ----------------
@app.route("/shipments")
def shipments():
    rows = Shipment.query.order_by(Shipment.id.desc()).all()
    status_counts = {}
    for s in rows:
        status_counts[s.status] = status_counts.get(s.status, 0) + 1
    status_breakdown = sorted(status_counts.items(), key=lambda x: -x[1])
    max_status = max([x[1] for x in status_breakdown]) if status_breakdown else 1
    total_containers = sum(s.container_count for s in rows)
    in_transit = len([s for s in rows if s.status in ("In Transit", "At Port", "Customs Clearance")])
    return render_template(
        "shipments.html", active="shipments", shipments=rows, total=len(rows),
        total_containers=total_containers, in_transit=in_transit,
        status_breakdown=status_breakdown, max_status=max_status,
    )


@app.route("/shipments/add", methods=["POST"])
def add_shipment():
    s = Shipment(
        shipment_ref=request.form["shipment_ref"],
        origin_country=request.form.get("origin_country", "China"),
        supplier_name=request.form.get("supplier_name", ""),
        items_description=request.form.get("items_description", ""),
        container_count=int(request.form.get("container_count") or 1),
        status=request.form.get("status", "Ordered"),
        port_of_entry=request.form.get("port_of_entry", ""),
        expected_arrival=request.form.get("expected_arrival", ""),
        actual_arrival=request.form.get("actual_arrival", ""),
        value=float(request.form.get("value") or 0),
        notes=request.form.get("notes", ""),
    )
    db.session.add(s)
    db.session.commit()
    flash("Shipment added successfully.")
    return redirect(url_for("shipments"))


@app.route("/shipments/edit/<int:id>", methods=["POST"])
def edit_shipment(id):
    s = Shipment.query.get_or_404(id)
    s.shipment_ref = request.form["shipment_ref"]
    s.origin_country = request.form.get("origin_country", "China")
    s.supplier_name = request.form.get("supplier_name", "")
    s.items_description = request.form.get("items_description", "")
    s.container_count = int(request.form.get("container_count") or 1)
    s.status = request.form.get("status", "Ordered")
    s.port_of_entry = request.form.get("port_of_entry", "")
    s.expected_arrival = request.form.get("expected_arrival", "")
    s.actual_arrival = request.form.get("actual_arrival", "")
    s.value = float(request.form.get("value") or 0)
    s.notes = request.form.get("notes", "")
    db.session.commit()
    flash("Shipment updated successfully.")
    return redirect(url_for("shipments"))


@app.route("/shipments/delete/<int:id>", methods=["POST"])
def delete_shipment(id):
    s = Shipment.query.get_or_404(id)
    db.session.delete(s)
    db.session.commit()
    return redirect(url_for("shipments"))


@app.route("/shipments/export")
def export_shipments():
    rows = Shipment.query.all()
    data = [[s.shipment_ref, s.origin_country, s.supplier_name, s.items_description,
             s.container_count, s.status, s.port_of_entry, s.expected_arrival, s.actual_arrival, s.value] for s in rows]
    return excel_response(
        ["Ref", "Origin", "Supplier", "Items", "Containers", "Status", "Port", "Expected Arrival", "Actual Arrival", "Value"],
        data, "shipments.xlsx", title="NEVF — Shipments Report"
    )


# ---------------- GLOBAL SEARCH ----------------
@app.route("/search")
def search():
    q = request.args.get("q", "").strip()
    results = []
    if q:
        like = f"%{q}%"

        for c in Customer.query.filter(or_(Customer.name.ilike(like), Customer.city.ilike(like))).limit(5).all():
            results.append({"type": "Customer", "label": c.name, "sub": c.city or "", "url": url_for("customers")})

        for l in Lead.query.filter(Lead.title.ilike(like)).limit(5).all():
            results.append({"type": "Lead", "label": l.title, "sub": l.stage, "url": url_for("leads")})

        for c in Contract.query.filter(or_(Contract.ref.ilike(like), Contract.customer_name.ilike(like))).limit(5).all():
            results.append({"type": "Contract", "label": c.ref, "sub": c.customer_name or "", "url": url_for("contracts")})

        for p in Plant.query.filter(or_(Plant.name.ilike(like), Plant.city.ilike(like))).limit(5).all():
            results.append({"type": "Plant", "label": p.name, "sub": p.city, "url": url_for("plants")})

        for e in Employee.query.filter(or_(Employee.name.ilike(like), Employee.designation.ilike(like))).limit(5).all():
            results.append({"type": "Employee", "label": e.name, "sub": e.designation or "", "url": url_for("employees")})

        for s in Shipment.query.filter(or_(Shipment.shipment_ref.ilike(like), Shipment.supplier_name.ilike(like))).limit(5).all():
            results.append({"type": "Shipment", "label": s.shipment_ref, "sub": s.supplier_name or "", "url": url_for("shipments")})

        for b in BusUnit.query.filter(BusUnit.unit_no.ilike(like)).limit(5).all():
            results.append({"type": "Bus Unit", "label": b.unit_no, "sub": b.model or "", "url": url_for("inventory_buses")})

        for c in CityFleet.query.filter(CityFleet.city.ilike(like)).limit(5).all():
            results.append({"type": "City", "label": c.city, "sub": f"{c.buses} buses", "url": url_for("fleet_city")})

    return {"results": results}


# ---------------- SETTINGS ----------------
@app.route("/settings")
def settings_page():
    s = get_settings()
    users = SystemUser.query.all()
    return render_template("settings.html", active="settings", settings=s, users=users)


@app.route("/settings/save", methods=["POST"])
def save_settings():
    s = get_settings()
    s.company_name = request.form.get("company_name", s.company_name)
    s.manufacturing_target = int(request.form.get("manufacturing_target") or s.manufacturing_target)
    s.currency = request.form.get("currency", s.currency)
    db.session.commit()
    flash("Settings saved.")
    return redirect(url_for("settings_page"))


@app.route("/settings/users/add", methods=["POST"])
def add_user():
    u = SystemUser(name=request.form["name"], email=request.form.get("email", ""), role=request.form.get("role", "Staff"))
    db.session.add(u)
    db.session.commit()
    return redirect(url_for("settings_page"))


@app.route("/settings/users/delete/<int:id>", methods=["POST"])
def delete_user(id):
    u = SystemUser.query.get_or_404(id)
    db.session.delete(u)
    db.session.commit()
    return redirect(url_for("settings_page"))


# ---------------- ENTRY POINT ----------------
if __name__ == "__main__":
    with app.app_context():
        db.create_all()
    app.run(debug=True, host="127.0.0.1", port=5000)